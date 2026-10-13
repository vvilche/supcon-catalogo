#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrae el texto de los documentos de SOFTWARE DE PROCESO y PLATAFORMA de SUPCON
(APC, RTO, EAM, LIMS, OMS, OTS, PLS, Blend, SCO, IFEM, PID, TPT, TAS, MP3000,
PRIDE, Nyx/UCS, ESP, supOS, supLink, OMC, ciberseguridad) desde
/Users/victorvilche/Documents/ConocimientoSupcon/Process - Supcon1..6 y Tier0.

Deduplica por MD5. Usa PyMuPDF para PDF y python-pptx para PPTX.
No usa OCR (basta la capa de texto). No inventa: si no hay texto, lo registra.

Salida: datos/software_extraido.json
"""
import os, sys, json, hashlib, re, warnings
warnings.filterwarnings("ignore")

RAIZ = "/Users/victorvilche/Documents/ConocimientoSupcon"
BASE = os.path.join(RAIZ, "supcon-catalogo")
OUT = os.path.join(BASE, "datos", "software_extraido.json")

# --- filtro de relevancia: software de proceso / plataforma / ciberseguridad ---
CLAVES = [
    "apc", "rto", "eam", "lims", "oms", "ots", "pls", "blend", "sco", "ifem",
    "pid", "tpt", "tas", "mp3000", "pride", "nyx", "ucs", "esp", "supos",
    "suplink", "omc", "cyber", "secur", "inplant", "defender", "q-lab", "qlab",
    "apex", "plantbot", "soli", "scada", "alarm", "workflow", "simulator",
    "operator training", "blending", "scheduling", "maint", "asset", "tank farm",
    "metering", "movimiento", "batch", "mes", "unified name", "uns",
]
EXCLUIR = [
    "certificat", "资质", "certificate", "_ce[", "ce-", "rohs", "eac", "atex",
    "iecex", "nepsi", "tuv", "sil3 certificate", "ul证书", "ce certificate",
]

def relevante(p):
    s = p.lower()
    nombre = os.path.basename(s)
    if any(x in s for x in EXCLUIR) and not any(
        k in nombre for k in ("apc", "rto", "eam", "oms", "ots", "tpt", "pride", "ucs", "supos", "lims", "blend", "sco", "pid", "tas")
    ):
        return False
    if "/analyser(hobre)" in s or "/instrument" in s or "/switchgear" in s:
        return False
    if "/03system hardware" in s or "/05past manual" in s or "/io module" in s:
        return False
    return any(k in nombre for k in CLAVES) or any(
        k in s for k in ("software", "sw/", "/sw", "miv business", "soliscada")
    )

def md5(path, bloque=1 << 20):
    h = hashlib.md5()
    with open(path, "rb") as f:
        while True:
            b = f.read(bloque)
            if not b:
                break
            h.update(b)
    return h.hexdigest()

def limpia(t):
    t = t.replace("\u00a0", " ").replace("\uf0b7", "-").replace("\u2022", "-")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()

def pdf_texto(path):
    import fitz
    doc = fitz.open(path)
    partes = []
    for i, pag in enumerate(doc):
        try:
            t = pag.get_text("text")
        except Exception as e:
            t = ""
            partes.append(f"[error pag {i+1}: {e}]")
        if t:
            partes.append(f"--- pag {i+1} ---\n{limpia(t)}")
    n = doc.page_count
    doc.close()
    return "\n".join(partes), n

def pptx_texto(path):
    from pptx import Presentation
    pr = Presentation(path)
    partes = []
    for i, slide in enumerate(pr.slides):
        lineas = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                lineas.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for fila in sh.table.rows:
                    celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                    if celdas:
                        lineas.append(" | ".join(celdas))
        if lineas:
            partes.append(f"--- slide {i+1} ---\n{limpia(chr(10).join(lineas))}")
    return "\n".join(partes), len(pr.slides._sldIdLst)

def docx_texto(path):
    from docx import Document
    d = Document(path)
    partes = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        for fila in tb.rows:
            c = [cel.text.strip() for cel in fila.cells if cel.text.strip()]
            if c:
                partes.append(" | ".join(c))
    return limpia("\n".join(partes)), len(d.paragraphs)

def main():
    cands = []
    for raiz, dirs, files in os.walk(RAIZ):
        if "supcon-catalogo" in raiz or ".git" in raiz:
            continue
        for fn in files:
            if fn.startswith("~$") or fn.startswith("."):
                continue
            ext = os.path.splitext(fn)[1].lower()
            if ext not in (".pdf", ".pptx", ".docx"):
                continue
            p = os.path.join(raiz, fn)
            if relevante(p):
                cands.append(p)
    print(f"candidatos: {len(cands)}")

    vistos = {}
    extraidos = []
    saltados = []
    for p in sorted(cands):
        try:
            h = md5(p)
        except Exception as e:
            saltados.append({"archivo": p, "motivo": f"md5: {e}"})
            continue
        if h in vistos:
            saltados.append({"archivo": p, "motivo": f"duplicado de {vistos[h]}", "md5": h})
            continue
        vistos[h] = p
        ext = os.path.splitext(p)[1].lower()
        try:
            if ext == ".pdf":
                txt, n = pdf_texto(p)
            elif ext == ".pptx":
                txt, n = pptx_texto(p)
            else:
                txt, n = docx_texto(p)
        except Exception as e:
            saltados.append({"archivo": p, "motivo": f"extract: {type(e).__name__}: {e}"})
            continue
        rel = os.path.relpath(p, RAIZ)
        extraidos.append({
            "archivo": rel,
            "tipo": ext.lstrip("."),
            "md5": h,
            "paginas_o_slides": n,
            "chars": len(txt),
            "texto": txt,
        })
        print(f"  OK {len(txt):7d} ch  {n:4d}p/s  {rel}")

    data = {"raiz": RAIZ, "n_documentos": len(extraidos), "documentos": extraidos,
            "descartados": saltados}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    print(f"\n=== {len(extraidos)} documentos unicos -> {OUT}")
    print(f"=== {len(saltados)} descartados por duplicado/error")
    vacios = [d["archivo"] for d in extraidos if d["chars"] < 200]
    if vacios:
        print(f"=== {len(vacios)} con poco/no texto (posible escaneo):")
        for v in vacios:
            print("    -", v)
    total = sum(d["chars"] for d in extraidos)
    print(f"=== total {total:,} caracteres extraidos")

if __name__ == "__main__":
    main()
